"""Optional upstream extraction: unchanged map PNGs and author-supplied workbook.

python3 extract_sources.py --maps /path/to/maps --workbook /path/to/workbook.xlsx --output extracted
Requires Pillow for maps and openpyxl for workbook reading. Neither input is modified.
The core reproduce.py analysis uses the supplied extracted inputs and standard library only.
"""
from pathlib import Path
from datetime import datetime, timedelta
import argparse, csv, hashlib, json, zipfile, xml.etree.ElementTree as ET

def write(path, rows):
    with path.open('w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)

def extract_maps(source, output):
    from PIL import Image
    # Equirectangular plot bounds calibrated against the supplied labelled map.
    lat = 5 + 24/60 + 49.116/3600
    lon = 118 + 2/60 + 16.112/3600
    x = round(31 + (lon + 180)/360*510)
    y = round(33 + (90-lat)/180*398)
    files = sorted(source.glob('RF_*.png'))
    if not files: raise ValueError('No RF_YYYYMMDD_HHMM.png maps found')
    patches, inventory, legend = [], [], None
    for path in files:
        im = Image.open(path).convert('RGB')
        if im.size != (640, 479): raise ValueError(f'Unexpected image size: {path.name}')
        this_legend = {v: im.getpixel((552, round(47+(14-v)*28.5))) for v in range(1,15)}
        if legend is None: legend = this_legend
        if legend != this_legend: raise ValueError(f'Legend differs: {path.name}')
        timestamp = datetime.strptime(path.stem, 'RF_%Y%m%d_%H%M').isoformat(' ')
        inventory.append({'filename':path.name, 'timestamp_UTC':timestamp,
                          'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
        for dx in range(-3,4):
            for dy in range(-3,4):
                rgb = im.getpixel((x+dx,y+dy))
                patches.append({'filename':path.name,'dx':dx,'dy':dy,
                                'red':rgb[0],'green':rgb[1],'blue':rgb[2]})
    write(output/'map_pixels.csv', patches)
    write(output/'map_inventory.csv', inventory)
    write(output/'map_palette.csv', [{'class_MHz':v,'red':c[0],'green':c[1],'blue':c[2]} for v,c in sorted(legend.items())])
    (output/'map_extraction.json').write_text(json.dumps({
        'latitude':lat,'longitude':lon,'central_pixel':[x,y], 'image_size':[640,479],
        'map_bounds_pixels':[31,33,541,431], 'patch':'7 x 7 original RGB pixels, containing nested 3 x 3 and 5 x 5 patches',
        'time_source':'Filename timestamps, interpreted as UTC; representative rendered titles checked during source review, not all titles independently OCR-verified',
        'classification':'Exact legend-colour matching; no nearest-colour interpolation',
        'maps':len(inventory),'pixels':len(patches)},indent=2)+'\n')

def extract_workbook(source, output):
    import openpyxl
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    for name, col in [('Darwin',10),('Guam',11)]:
        sheet=wb[name]
        if str(sheet.cell(20,col).value) != '2023': raise ValueError('Unexpected year header')
        rows=[]
        for row in sheet.iter_rows(min_row=21,max_row=1364,max_col=col,values_only=True):
            day, clock, value = row[0],row[1],row[col-1]
            t=datetime.combine(day.date().replace(year=2023),clock)
            rows.append({'April_day':15+(t-datetime(2023,4,15)).total_seconds()/86400,
                         'foF2_MHz':value})
        write(output/f'{name}_2023_foF2.csv',rows)
    wb.close()

def extract_power_charts(source, output):
    ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    settings=[('chart1.xml',5.357,'2023-04-18'),('chart3.xml',7.078,'2023-04-20'),('chart4.xml',10.130,'2023-04-23')]
    points=[]
    with zipfile.ZipFile(source) as z:
        for name,frequency,start in settings:
            root=ET.fromstring(z.read('word/charts/'+name));series=root.find('.//c:ser',ns)
            def values(field):
                pts=series.findall(f'c:{field}//c:pt',ns)
                return [float(p.find('c:v',ns).text) for p in sorted(pts,key=lambda p:int(p.attrib['idx']))]
            xs,ys=values('cat'),values('val');require_equal=len(xs)==len(ys)
            if not require_equal:raise ValueError('Chart coordinate lengths differ')
            base=datetime.fromisoformat(start);previous=-1
            for i,(x,y) in enumerate(zip(xs,ys)):
                seconds=round(x*86400)
                if seconds<previous:base+=timedelta(days=1)
                previous=seconds
                points.append({'chart':name,'point_index':i,'recorded_timestamp':str(base+timedelta(seconds=seconds)),
                               'frequency_MHz':frequency,'SNR_dB':y,'nominal_power_W':1})
    write(output/'nominal_power_chart_points.csv',points)

if __name__ == '__main__':
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--maps',type=Path);ap.add_argument('--workbook',type=Path);ap.add_argument('--trial-report',type=Path)
    ap.add_argument('--output',required=True,type=Path);args=ap.parse_args()
    args.output.mkdir(parents=True,exist_ok=True)
    if args.maps:extract_maps(args.maps,args.output)
    if args.workbook:extract_workbook(args.workbook,args.output)
    if args.trial_report:extract_power_charts(args.trial_report,args.output)
